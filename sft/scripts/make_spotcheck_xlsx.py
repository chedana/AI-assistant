"""阶段 4 人工抽检表：每个测试集随机抽 15%（种子 15），导出给负责人逐条核对。"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
rows = json.loads((ROOT / "data/test_v1/spotcheck_sample.json").read_text())
FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="1F3864")
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def gold_text(g):
    c = g["constraints"]; parts = []
    if c["location_keywords"]: parts.append("地点 " + " / ".join(c["location_keywords"]))
    if c["commute_destination"]: parts.append("通勤 " + c["commute_destination"])
    if c["max_rent"]: parts.append(f"预算 ≤ {c['max_rent']['value']:,} {c['max_rent']['period']}")
    for o in c["layout_options"]:
        bits = []
        if o["bedrooms"] is not None: bits.append(f"{o['bedrooms']} 卧")
        if o["bathrooms"] is not None: bits.append(f"{o['bathrooms']} 卫")
        if o["property_type"]: bits.append(o["property_type"])
        parts.append("户型 " + " ".join(bits))
    for k, label in (("furnish_type", "家具"), ("let_type", "租期"), ("available_from", "入住")):
        if c[k]: parts.append(f"{label} {c[k]}")
    names = {"pets_allowed": "宠物", "garden": "花园", "parking": "停车", "bills_included": "含账单",
             "student_friendly": "接受学生", "families_allowed": "接受家庭", "smokers_allowed": "接受吸烟",
             "dss_accepted": "接受 DSS"}
    for k, label in names.items():
        if c[k] is not None: parts.append(f"{label}={'是' if c[k] else '否'}")
    return "；".join(parts) if parts else "（全部为空）"


wb = Workbook()
ws = wb.active
ws.title = "抽检"
legend = [
    "测试集人工抽检（test_v1）：每个测试集随机抽 15%，共 90 条。",
    "只需填写黄色的两列：「判定」从下拉框选 对 / 错；判定为错时在「错在哪」写明原因，例如：错｜预算周期应为 pw 而不是 pcm。",
    "判断标准：gold 是否准确反映 query 的意思（意图、硬条件不多不少、地点原样、软偏好是可选喜好）；query 是否像真实用户会说的英文。",
]
for i, t in enumerate(legend, 1):
    ws.cell(row=i, column=1, value=t).font = Font(name=FONT, bold=(i == 1), size=12 if i == 1 else 10)
HDR = 5
cols = ["序号", "测试集", "ID", "类型", "生成方", "上一轮展示的房源", "用户 query", "gold 意图",
        "gold 硬条件", "gold 软偏好", "风格标签", "判定（对/错）", "错在哪"]
widths = [6, 16, 10, 12, 8, 42, 48, 12, 46, 30, 22, 12, 36]
for j, (h, w) in enumerate(zip(cols, widths), 1):
    c = ws.cell(row=HDR, column=j, value=h)
    c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = HEAD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BORDER
    ws.column_dimensions[c.column_letter].width = w
kind_cn = {"search": "搜索", "followup": "追问", "softpref": "软偏好", "specific_qa": "问某一套",
           "compare": "比较", "chitchat": "闲聊"}
for i, r in enumerate(rows, 1):
    rr = HDR + i
    ctx = "\n".join(r["context"].split("\n")[1:]) if r["context"] else ""
    vals = [i, "抽取" if r["set"] == "test_extract_v1" else "意图", r["id"], kind_cn[r["kind"]], r["generator"],
            ctx, r["query"], r["gold"]["intent"], gold_text(r["gold"]),
            " / ".join(r["gold"]["soft_preferences"]), ", ".join(r["style_tags"]), None, None]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=rr, column=j, value=v)
        c.font = Font(name=FONT, size=10); c.border = BORDER
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if j >= 12: c.fill = YELLOW
last = HDR + len(rows)
dv = DataValidation(type="list", formula1='"对,错"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"L{HDR+1}:L{last}")
ws.freeze_panes = ws.cell(row=HDR + 1, column=4)
ws.auto_filter.ref = f"A{HDR}:M{last}"

s = wb.create_sheet("汇总")
s["A1"] = "抽检汇总（自动计算，不用填写）"; s["A1"].font = Font(name=FONT, bold=True, size=12)
heads = ["测试集", "抽检条数", "已判定", "判对", "判错", "错误率（已判定中）"]
for j, h in enumerate(heads, 1):
    c = s.cell(row=3, column=j, value=h); c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = HEAD
    s.column_dimensions[c.column_letter].width = 18
rng_set, rng_res = f"抽检!$B${HDR+1}:$B${last}", f"抽检!$L${HDR+1}:$L${last}"
for k, name in enumerate(["抽取", "意图"]):
    r = 4 + k
    s.cell(row=r, column=1, value=name)
    s.cell(row=r, column=2, value=f'=COUNTIF({rng_set},A{r})')
    s.cell(row=r, column=3, value=f'=COUNTIFS({rng_set},A{r},{rng_res},"对")+COUNTIFS({rng_set},A{r},{rng_res},"错")')
    s.cell(row=r, column=4, value=f'=COUNTIFS({rng_set},A{r},{rng_res},"对")')
    s.cell(row=r, column=5, value=f'=COUNTIFS({rng_set},A{r},{rng_res},"错")')
    s.cell(row=r, column=6, value=f'=IF(C{r}=0,"-",E{r}/C{r})')
s.cell(row=6, column=1, value="合计")
for col in "BCDE":
    s[f"{col}6"] = f"=SUM({col}4:{col}5)"
s["F6"] = '=IF(C6=0,"-",E6/C6)'
for r in range(4, 7):
    s[f"F{r}"].number_format = "0.0%"
    for col in "ABCDEF":
        s[f"{col}{r}"].font = Font(name=FONT, bold=(r == 6))
s["A8"] = "抽检样本：data/test_v1/spotcheck_sample.json（种子 15，每个测试集 15%）。错误率过高时，修正后另存测试集 v2，不在 v1 上直接改。"
s["A8"].font = Font(name=FONT, size=9, italic=True)
out = ROOT / "results/spotcheck_test_v1.xlsx"
wb.save(out); print(out)
